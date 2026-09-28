"""Read XLS/XLSX into the same cell/merge model, without executing formulas."""

from dataclasses import dataclass
from io import BytesIO, StringIO
from zipfile import ZipFile

import openpyxl
import xlrd

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 100


class ParseError(ValueError):
    """An unsupported or invalid workbook, safe to explain to the user."""


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


@dataclass
class Sheet:
    name: str
    cells: dict[tuple[int, int], str]
    merges: list[tuple[int, int, int, int]]  # r0, r1, c0, c1; exclusive ends
    height: int
    width: int

    def text(self, row: int, column: int) -> str:
        # Deliberately do not fill down merged cells into another lesson.
        return self.cells.get((row, column), "")

    def merge_at(self, row: int, column: int):
        return next((m for m in self.merges if m[0] <= row < m[1] and m[2] <= column < m[3]), None)


def read_workbook(data: bytes) -> list[Sheet]:
    if not data or len(data) > MAX_FILE_BYTES:
        raise ParseError("Нужен непустой Excel размером до 5 МБ.")
    try:
        if data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            book = xlrd.open_workbook(file_contents=data, formatting_info=True, logfile=StringIO())
            try:
                sheets = []
                for source in book.sheets():
                    cells = {}
                    if source.nrows <= MAX_ROWS and source.ncols <= MAX_COLUMNS:
                        cells = {(r, c): clean(source.cell_value(r, c)) for r in range(source.nrows)
                                 for c in range(source.ncols) if source.cell_value(r, c) != ""}
                    sheets.append(Sheet(source.name, cells, source.merged_cells, source.nrows, source.ncols))
                return sheets
            finally:
                book.release_resources()
        if data.startswith(b"PK"):
            with ZipFile(BytesIO(data)) as archive:
                if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                    raise ParseError("Распакованный Excel слишком большой для расписания одной группы.")
            book = openpyxl.load_workbook(BytesIO(data), data_only=True, keep_links=False)
            try:
                sheets = []
                for source in book.worksheets:
                    cells = {}
                    if source.max_row <= MAX_ROWS and source.max_column <= MAX_COLUMNS:
                        cells = {(cell.row - 1, cell.column - 1): clean(cell.value) for row in source
                                 for cell in row if cell.value is not None}
                    merges = [(m.min_row - 1, m.max_row, m.min_col - 1, m.max_col) for m in source.merged_cells.ranges]
                    sheets.append(Sheet(source.title, cells, merges, source.max_row, source.max_column))
                return sheets
            finally:
                book.close()
    except ParseError:
        raise
    except Exception as error:
        raise ParseError("Не удалось открыть Excel. Файл повреждён, защищён паролем или имеет неподдерживаемый формат.") from error
    raise ParseError("Это не двоичный Excel .xls/.xlsx. Переименование расширения не поможет.")
